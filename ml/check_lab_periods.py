from collections import defaultdict
from openpyxl import load_workbook

from backend.database import query

WORKBOOK = (
    "training_data/venue_data/"
    "LabPulse_AI_Lab_Sessions_Formatted_With_Venues.xlsx"
)


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_year(value):
    text = clean(value).upper()

    mapping = {
        "IST YEAR": "IST YEAR",
        "IST": "IST YEAR",
        "1ST YEAR": "IST YEAR",
        "FIRST YEAR": "IST YEAR",
        "2ND YEAR": "2ND",
        "SECOND YEAR": "2ND",
        "3RD YEAR": "3RD",
        "THIRD YEAR": "3RD",
        "4TH YEAR": "4TH",
        "FOURTH YEAR": "4TH",
    }

    return mapping.get(text, text)


def period_number(value):
    text = clean(value).upper()
    if text.startswith("PERIOD-"):
        try:
            return int(text.split("-")[1])
        except (ValueError, IndexError):
            return None
    return None


def normalize_subject(value):
    return clean(value).upper()


def main():
    rows = query("""
        SELECT academic_year, department, section, day,
               period, subject, time_slot
        FROM timetable
        WHERE resource_type = 'LAB'
    """)

    timetable_groups = defaultdict(list)

    for row in rows:
        key = (
            row["academic_year"],
            row["department"],
            row["section"],
            row["day"],
            normalize_subject(row["subject"]),
        )
        number = period_number(row["period"])

        if number is not None:
            timetable_groups[key].append(number)

    workbook = load_workbook(
        WORKBOOK,
        read_only=True,
        data_only=True,
    )
    sheet = workbook["Lab Sessions"]

    headers = [clean(c.value) for c in sheet[3]]
    columns = {
        name: i for i, name in enumerate(headers) if name
    }

    matched = []
    missing = []
    ambiguous = []
    blank_venues = 0

    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row[columns["Academic Year"]]:
            continue

        year = normalize_year(row[columns["Academic Year"]])
        department = clean(row[columns["Department"]]).upper()
        section = clean(row[columns["Section"]]).upper()
        day = clean(row[columns["Day"]]).upper()
        subject = normalize_subject(row[columns["Lab Name"]])
        venue = clean(row[columns["Lab Venue"]])

        start = period_number(row[columns["Start Period"]])
        end = period_number(row[columns["End Period"]])

        if not venue:
            blank_venues += 1

        key = (year, department, section, day, subject)
        actual_periods = sorted(set(timetable_groups.get(key, [])))

        if start is None or end is None or end < start:
            ambiguous.append((key, "Invalid period range"))
            continue

        expected = list(range(start, end + 1))
        matching = [
            p for p in actual_periods
            if start <= p <= end
        ]

        if matching == expected:
            matched.append({
                "key": key,
                "periods": expected,
                "venue": venue or "(blank)",
            })
        elif not matching:
            missing.append({
                "key": key,
                "expected": expected,
                "available": actual_periods,
            })
        else:
            ambiguous.append({
                "key": key,
                "expected": expected,
                "available_in_range": matching,
                "all_available": actual_periods,
            })

    workbook.close()

    print("TIMETABLE LAB ROWS:", len(rows))
    print("GROUPED SESSIONS CHECKED:", len(matched) + len(missing) + len(ambiguous))
    print("EXACT PERIOD MATCHES:", len(matched))
    print("MISSING / NO PERIOD MATCH:", len(missing))
    print("PARTIAL / AMBIGUOUS MATCHES:", len(ambiguous))
    print("BLANK VENUE SESSIONS:", blank_venues)

    print("\nSAMPLE EXACT MATCHES:")
    for item in matched[:10]:
        print(item)

    print("\nSAMPLE MISSING SESSIONS:")
    for item in missing[:10]:
        print(item)

    print("\nSAMPLE AMBIGUOUS SESSIONS:")
    for item in ambiguous[:10]:
        print(item)


if __name__ == "__main__":
    main()