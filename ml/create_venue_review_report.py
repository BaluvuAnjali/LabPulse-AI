from openpyxl import load_workbook
import csv
from pathlib import Path

base = Path("training_data/venue_data")
source = base / "LabPulse_AI_Lab_Sessions_Formatted_With_Venues.xlsx"
output = base / "venue_review_report.csv"

wb = load_workbook(source, data_only=True, read_only=True)
ws = wb["Lab Sessions"]

headers = [ws.cell(3, col).value for col in range(1, 15)]
rows = []

for values in ws.iter_rows(min_row=4, max_col=14, values_only=True):
    record = dict(zip(headers, values))
    venue = str(record.get("Lab Venue") or "").strip()
    venue_type = (
        "LABORATORY" if venue and "LAB" in venue.upper()
        else "OTHER" if venue
        else "BLANK"
    )

    rows.append({
        "Academic Year": record.get("Academic Year"),
        "Department": record.get("Department"),
        "Section": record.get("Section"),
        "Day": record.get("Day"),
        "Lab Name": record.get("Lab Name"),
        "Lab Venue": venue,
        "Venue Classification (initial)": venue_type,
        "Start Period": record.get("Start Period"),
        "End Period": record.get("End Period"),
        "Time": record.get("Time"),
        "Faculty": record.get("Faculty"),
        "Students": record.get("Students"),
        "Review Status": "REVIEW_REQUIRED",
        "Physical Identity Verified": "NO",
        "Availability Verified": "NO",
    })

wb.close()

with output.open("w", newline="", encoding="utf-8-sig") as file:
    writer = csv.DictWriter(file, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)

print(f"Created: {output}")
print(f"Sessions included: {len(rows)}")
print(f"Blank venue assignments: {sum(not r['Lab Venue'] for r in rows)}")
print("Database changes: NONE")
