from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

base = Path("training_data/venue_data")
source = base / "LabPulse_AI_Lab_Sessions_Formatted_With_Venues.xlsx"
output = base / "LabPulse_Venue_Verification.xlsx"

src = load_workbook(source, data_only=True, read_only=True)
ws = src["Lab Sessions"]

venues = {}
for row in ws.iter_rows(min_row=4, max_col=14, values_only=True):
    name = str(row[5] or "").strip()
    if not name:
        continue
    venues[name] = venues.get(name, 0) + 1
src.close()

wb = Workbook()
sheet = wb.active
sheet.title = "Venue Verification"

headers = [
    "Venue Name from Timetable",
    "Mapped Sessions",
    "Provisional Type",
    "Physical Venue Name Confirmed",
    "Separate Physical Venue?",
    "Location / Block",
    "Capacity (official only)",
    "Approved by",
    "Verification Status",
    "Notes",
]
sheet.append(headers)

for name, count in sorted(venues.items()):
    if "+" in name:
        kind = "Combined name - review"
    elif "SEMINAR HALL" in name.upper():
        kind = "Seminar hall - review"
    elif "CLASS ROOM" in name.upper():
        kind = "Classroom - review"
    elif "LAB" in name.upper():
        kind = "Laboratory candidate - verify"
    else:
        kind = "Other - review"

    sheet.append([
        name, count, kind, "Pending", "Pending",
        "", "", "", "REVIEW_REQUIRED", ""
    ])

for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="244062")
    cell.alignment = Alignment(wrap_text=True, vertical="center")

widths = [38, 16, 30, 28, 24, 24, 22, 22, 22, 36]
for index, width in enumerate(widths, start=1):
    sheet.column_dimensions[get_column_letter(index)].width = width

sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
sheet.row_dimensions[1].height = 35

wb.save(output)
print("Created:", output)
print("Distinct nonblank venue names:", len(venues))
print("Database changes: NONE")
