from openpyxl import load_workbook
from backend.database import get_db

WORKBOOK = (
    "training_data/venue_data/"
    "LabPulse_AI_Lab_Sessions_Formatted_With_Venues.xlsx"
)

SOURCE_FILE = "LabPulse_AI_Lab_Sessions_Formatted_With_Venues.xlsx"


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_year(value):
    text = clean(value).upper()

    mapping = {
        "IST": "IST YEAR",
        "IST YEAR": "IST YEAR",
        "1ST YEAR": "IST YEAR",
        "FIRST YEAR": "IST YEAR",
        "2ND YEAR": "2ND",
        "SECOND YEAR": "2ND",
        "3RD YEAR": "3RD",
        "THIRD YEAR": "3RD",
        "4TH YEAR": "4TH",
        "FOURTH YEAR": "4TH",
    }

    return mapping.get(text, text)


def period_number(value):
    text = clean(value).upper()

    if not text.startswith("PERIOD-"):
        raise ValueError(f"Unexpected period value: {value!r}")

    return int(text.split("-", 1)[1])


def classify_venue(venue):
    text = venue.upper()

    if not text:
        return "UNSPECIFIED"
    if "CLASS ROOM" in text or "CLASSROOM" in text:
        return "CLASSROOM"
    if "SEMINAR HALL" in text:
        return "SEMINAR_HALL"
    if "LAB" in text:
        return "LABORATORY"

    return "OTHER"


def main():
    workbook = load_workbook(
        WORKBOOK,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Lab Sessions"]
        headers = [clean(cell.value) for cell in sheet[3]]
        column = {
            name: index
            for index, name in enumerate(headers)
            if name
        }

        required = [
            "Academic Year",
            "Department",
            "Section",
            "Day",
            "Lab Name",
            "Lab Venue",
            "Start Period",
            "End Period",
        ]

        missing = [name for name in required if name not in column]
        if missing:
            raise ValueError(
                f"Workbook is missing required columns: {missing}"
            )

        imported = 0
        skipped_blank = 0
        rows_to_import = []
        seen = set()

        for row in sheet.iter_rows(min_row=4, values_only=True):
            year = normalize_year(row[column["Academic Year"]])
            department = clean(row[column["Department"]]).upper()
            section = clean(row[column["Section"]]).upper()
            day = clean(row[column["Day"]]).upper()
            subject = clean(row[column["Lab Name"]]).upper()
            venue = clean(row[column["Lab Venue"]])
            start_period = period_number(row[column["Start Period"]])
            end_period = period_number(row[column["End Period"]])

            # Venue assignments that are blank remain blank and are not imported.
            if not venue:
                skipped_blank += 1
                continue

            key = (
                year,
                department,
                section,
                subject,
                day,
                start_period,
                end_period,
            )

            if key in seen:
                raise ValueError(
                    f"Duplicate session key in workbook: {key}"
                )

            seen.add(key)

            rows_to_import.append((
                year,
                department,
                section,
                subject,
                day,
                f"Period-{start_period}..Period-{end_period}",
                venue,
                classify_venue(venue),
                "REVIEW_REQUIRED",
                SOURCE_FILE,
                (
                    "Imported from formatted lab-session workbook. "
                    "Venue assignment requires verification."
                ),
                start_period,
                end_period,
            ))

        with get_db() as db:
            for record in rows_to_import:
                (
                    year,
                    department,
                    section,
                    subject,
                    day,
                    time_slot,
                    venue,
                    venue_type,
                    status,
                    source_file,
                    notes,
                    start_period,
                    end_period,
                ) = record

                db.execute("""
                    INSERT INTO venue_mappings (
                        academic_year,
                        department,
                        section,
                        subject,
                        day,
                        time_slot,
                        lab_venue,
                        venue_type,
                        mapping_status,
                        source_file,
                        notes,
                        start_period,
                        end_period
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT (
                        academic_year,
                        department,
                        section,
                        subject,
                        day,
                        time_slot
                    )
                    DO UPDATE SET
                        lab_venue = excluded.lab_venue,
                        venue_type = excluded.venue_type,
                        mapping_status = excluded.mapping_status,
                        source_file = excluded.source_file,
                        notes = excluded.notes,
                        start_period = excluded.start_period,
                        end_period = excluded.end_period
                """, record)

                imported += 1

        print("VENUE IMPORT COMPLETED")
        print(f"Nonblank venue mappings imported: {imported}")
        print(f"Blank venue sessions preserved as blank: {skipped_blank}")
        print("Mapping status: REVIEW_REQUIRED")
        print("Timetable table: not modified")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()